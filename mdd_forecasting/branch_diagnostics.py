"""Diagnostyka wyłącznie odłożonego testu i raport Excel dla użytkownika."""
from __future__ import annotations

from pathlib import Path
import json

import numpy as np
import pandas as pd

from .config import WEATHER_FEATURES


def error_statistics(group: pd.DataFrame) -> dict:
    y = group["realizacja"].to_numpy(dtype=float, na_value=np.nan)
    p = group["prognoza"].to_numpy(dtype=float, na_value=np.nan)
    b = group["profil"].to_numpy(dtype=float, na_value=np.nan)
    valid = np.isfinite(y) & np.isfinite(p)
    paired = valid & np.isfinite(b)
    actual, prediction = y[valid], p[valid]
    error = prediction - actual
    absolute = np.abs(error)
    denominator = np.abs(actual).sum()
    mae = float(absolute.mean()) if len(error) else np.nan
    pair_mae = float(np.abs(p[paired] - y[paired]).mean()) if paired.any() else np.nan
    base_mae = float(np.abs(b[paired] - y[paired]).mean()) if paired.any() else np.nan
    smape_denominator = np.abs(actual) + np.abs(prediction)
    smape = np.divide(2 * absolute, smape_denominator, out=np.zeros_like(absolute), where=smape_denominator > 0)
    total_variance = ((actual - actual.mean()) ** 2).sum() if len(actual) else 0
    return {
        "n": int(valid.sum()), "n_zlecen": len(group),
        "n_godzin": group.loc[valid, "target_time_utc"].nunique(),
        "pokrycie": float(valid.mean()), "n_wspolne_z_profilem": int(paired.sum()),
        "mae": mae, "rmse": float(np.sqrt(np.mean(error ** 2))) if len(error) else np.nan,
        "bias": float(error.mean()) if len(error) else np.nan,
        "mediana_bledu_abs": float(np.median(absolute)) if len(error) else np.nan,
        "p90_bledu_abs": float(np.quantile(absolute, .9)) if len(error) else np.nan,
        "p95_bledu_abs": float(np.quantile(absolute, .95)) if len(error) else np.nan,
        "max_bledu_abs": float(absolute.max()) if len(error) else np.nan,
        "wape": float(absolute.sum() / denominator) if denominator > 0 else np.nan,
        "smape": float(smape.mean()) if len(error) else np.nan,
        "r2": float(1 - (error ** 2).sum() / total_variance) if total_variance > 0 else np.nan,
        "mae_ml_wspolne": pair_mae, "mae_profil_wspolne": base_mae,
        "poprawa_mae_vs_profil": float(1 - pair_mae / base_mae) if base_mae > 0 else np.nan,
    }


def aggregate_errors(frame: pd.DataFrame, keys: list[str]) -> pd.DataFrame:
    rows = []
    for values, group in frame.groupby(keys, dropna=False, observed=True, sort=True):
        if not isinstance(values, tuple):
            values = (values,)
        rows.append({**dict(zip(keys, values)), **error_statistics(group)})
    return pd.DataFrame(rows)


def diagnostic_tables(details: pd.DataFrame, balance: pd.DataFrame, timezone: str) -> dict[str, pd.DataFrame]:
    keys = ["oddzial_code", "forecast_origin_utc", "target_time_utc", "horyzont_h"]
    direction = details[[*keys, "kierunek_energii_norm", "wartosc_rzeczywista",
                         "wartosc_przewidywana", "wartosc_bazowa", *WEATHER_FEATURES]].rename(columns={
        "kierunek_energii_norm": "kierunek", "wartosc_rzeczywista": "realizacja",
        "wartosc_przewidywana": "prognoza", "wartosc_bazowa": "profil"})
    net = balance[[*keys, "bilans_wartosc_rzeczywista", "bilans_wartosc_przewidywana", "bilans_wartosc_bazowa"]].rename(columns={
        "bilans_wartosc_rzeczywista": "realizacja", "bilans_wartosc_przewidywana": "prognoza",
        "bilans_wartosc_bazowa": "profil"})
    net["kierunek"] = "BILANS"
    weather = details[[*keys, *WEATHER_FEATURES]].drop_duplicates(keys)
    net = net.merge(weather, on=keys, how="left", validate="one_to_one")
    errors = pd.concat([direction, net], ignore_index=True)
    local = errors["target_time_utc"].dt.tz_convert(timezone)
    errors["godzina_lokalna"] = local.dt.hour
    errors["dzien_tygodnia"] = local.dt.dayofweek
    errors["miesiac"] = local.dt.month
    errors["doba_lokalna"] = local.dt.tz_localize(None).dt.normalize()
    errors["blad"] = errors["prognoza"] - errors["realizacja"]
    errors["blad_abs"] = errors["blad"].abs()
    summary = aggregate_errors(errors, ["oddzial_code", "kierunek"])
    summary["ocena_vs_profil"] = np.select(
        [summary["n_wspolne_z_profilem"].eq(0), summary["mae_ml_wspolne"].gt(summary["mae_profil_wspolne"]),
         summary["mae_ml_wspolne"].lt(summary["mae_profil_wspolne"])],
        ["BRAK_WSPOLNEJ_OCENY", "ML_GORSZY_OD_PROFILU", "ML_LEPSZY_OD_PROFILU"], default="REMIS")
    tables = {"bledy_oddzialy": summary}
    for name, dimension in [("bledy_godziny", "godzina_lokalna"), ("bledy_dni_tygodnia", "dzien_tygodnia"),
                            ("bledy_miesiace", "miesiac"), ("bledy_doby", "doba_lokalna"),
                            ("bledy_horyzont", "horyzont_h")]:
        tables[name] = aggregate_errors(errors, ["oddzial_code", "kierunek", dimension])
    tables["bledy_przebiegi"] = aggregate_errors(errors, ["oddzial_code", "kierunek", "forecast_origin_utc"])
    tables["najwieksze_bledy"] = (errors.loc[np.isfinite(errors["blad_abs"])]
                                  .sort_values("blad_abs", ascending=False)
                                  .groupby(["oddzial_code", "kierunek"], sort=False).head(20))
    condition_tables = []
    for column, cuts in [("temperatura", [-np.inf, 0, 10, 20, 30, np.inf]),
                          ("promieniowanie_calkowite", [-np.inf, 0, 100, 300, 600, np.inf]),
                          ("predkosc_wiatru", [-np.inf, 2, 5, 10, np.inf])]:
        subset = errors.copy()
        subset["przedzial"] = pd.cut(subset[column], bins=cuts).astype("string").fillna("BRAK_DANYCH")
        table = aggregate_errors(subset, ["oddzial_code", "kierunek", "przedzial"])
        table.insert(2, "zmienna", column)
        condition_tables.append(table)
    tables["bledy_warunki_pogody"] = pd.concat(condition_tables, ignore_index=True)
    # Rozkłady są opisem testu. Nie wracają do wyboru modelu ani parametrów.
    return tables


def weather_coverage(joined: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for (branch, direction), group in joined.groupby(["oddzial_code", "kierunek_energii_norm"]):
        row = {"oddzial_code": branch, "kierunek": direction, "n_zlecen": len(group),
               "udzial_pogody_asof": float(group["pogoda_dopasowana"].mean()),
               "udzial_realizacji": float(group["wartosc_rzeczywista"].notna().mean())}
        row.update({f"pokrycie_{col}": float(group[col].notna().mean()) for col in WEATHER_FEATURES})
        rows.append(row)
    return pd.DataFrame(rows)


READ_ME = [
    ("Ocena", "Statystyki pochodzą wyłącznie z odłożonego, późniejszego testu. Końcowy zapisany model jest następnie uczony także na tym okresie."),
    ("Wybór modelu", "Rodzinę modelu, pogodę i parametry wybiera chronologiczna walidacja. Nie dobieraj ich ponownie po obejrzeniu tego samego testu."),
    ("MAE / RMSE", "Błąd w jednostce źródłowej energii. Mniej = lepiej. RMSE silniej karze duże pomyłki."),
    ("Bias", "Średnia prognoza minus realizacja. Dodatni = zawyżanie; ujemny = zaniżanie."),
    ("P90 / P95", "90% / 95% ocenionych błędów bezwzględnych nie przekracza tej wartości. To opis testu, nie przedział prognozy."),
    ("WAPE", "Suma błędów bezwzględnych / suma bezwzględnych realizacji; 0,10 = 10%. Przy małym bilansie bywa niestabilny."),
    ("sMAPE", "Symetryczny błąd względny (0–200%). Para zero/zero daje 0. Nie używamy MAPE dzielącego przez zerowe oddanie."),
    ("R2", "Wynik względem średniej realizacji; może być ujemny. Brak wartości przy stałym targetcie."),
    ("Porównanie z profilem", "Poprawa MAE jest liczona na wspólnych wierszach ML i profilu. Dodatnia = ML lepszy. Gdy MAE profilu wynosi 0, procent jest pusty."),
    ("Liczność", "n to pary origin/cel, więc ta sama godzina może wystąpić przy różnych horyzontach. n_godzin to liczba unikalnych godzin UTC."),
    ("Pokrycie oceny", "Statystyki dotyczą dopasowanej historii z realizacją. Pokrycie całej historii pogodą jest osobno w Pokrycie_pogody; nie traktuj n jako wszystkich godzin eksportu."),
    ("Pogoda", "Ważność permutacyjna grup jest opisem wpływu na błąd testu, nie dowodem przyczynowości. Ujemna wartość może oznaczać szum; nieużyta grupa nie oznacza udowodnionej bezużyteczności."),
    ("Przedziały pogody", "Przedziały w Warunki_pogody są w jednostkach źródłowych. Potwierdź jednostki z hurtownią przed interpretacją."),
    ("Czas", "Cele i origins są UTC; godzina, dzień i miesiąc diagnostyki są w Europe/Warsaw. Dwie jesienne 02:00 mają różne UTC."),
    ("Korekty", "Ocena retrospektywna na najlepszych dostępnych korektach; brak historii ich publikacji. Lagi energii pozostają wyłączone."),
    ("Klienci", "Obecność MDD nie jest potwierdzeniem aktywności umowy. Braki i zmiany portfela trzeba wyjaśnić."),
    ("Duże tabele", "Excel zawiera do 20 000 wierszy każdej tabeli. Pełne tabele są w plikach CSV obok raportu."),
]


def write_training_workbook(bundle: dict, reports: dict[str, pd.DataFrame], destination: str | Path,
                            *, read_me=None):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.utils import get_column_letter
    from .report import _excel_scalar

    workbook = Workbook()
    workbook.remove(workbook.active)

    def add_sheet(name: str, frame: pd.DataFrame):
        sheet = workbook.create_sheet(name[:31])
        sheet.append([str(c) for c in frame.columns])
        for row_number, values in enumerate(frame.head(20_000).itertuples(index=False, name=None), start=2):
            cleaned = []
            for value in values:
                if isinstance(value, (list, dict, tuple)):
                    value = json.dumps(value, ensure_ascii=False, default=str)
                value = _excel_scalar(value)
                cleaned.append(value)
            sheet.append(cleaned)
            # Nazwy i opisy źródłowe zachowujemy jako tekst, także przy prefiksie '='.
            for column_number, value in enumerate(cleaned, start=1):
                if isinstance(value, str):
                    sheet.cell(row=row_number, column=column_number).data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 32
        for position, column in enumerate(frame.columns, 1):
            sample = frame[column].head(60).map(str)
            width = min(46, max(13, len(str(column)) + 2, sample.str.len().max() + 1 if len(sample) else 0))
            sheet.column_dimensions[get_column_letter(position)].width = width
            if any(word in str(column) for word in ["wape", "smape", "pokrycie", "poprawa", "udzial"]):
                number_format = "0.0%"
            elif pd.api.types.is_numeric_dtype(frame[column]):
                number_format = "#,##0.000"
            else:
                number_format = None
            if number_format:
                for cells in sheet.iter_rows(min_row=2, min_col=position, max_col=position):
                    cells[0].number_format = number_format
        return sheet

    intro = add_sheet("Czytaj_najpierw", pd.DataFrame(READ_ME if read_me is None else read_me, columns=["temat", "wyjasnienie"]))
    intro.column_dimensions["A"].width = 29
    intro.column_dimensions["B"].width = 115
    for row in intro.iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
        intro.row_dimensions[row[0].row].height = 44
    settings_rows = [{"parametr": key, "wartosc": value} for key, value in bundle["settings"].items()]
    add_sheet("Ustawienia", pd.DataFrame(settings_rows))
    chosen = [{"kierunek": direction, "pogoda": spec["wariant"],
               "model": spec["parameters"].get("backend", "hist_gradient_boosting"),
               "parametry": spec["parameters"], "cechy": spec["columns"]}
              for direction, spec in bundle["models"].items()]
    add_sheet("Wybrane_modele", pd.DataFrame(chosen))
    sheets = {
        "bledy_oddzialy": "Podsumowanie", "metryki_test": "Horyzonty_24h",
        "bledy_godziny": "Godziny", "bledy_dni_tygodnia": "Dni_tygodnia", "bledy_miesiace": "Miesiace",
        "bledy_horyzont": f"Horyzont_1_do_{bundle['settings']['horizon_hours']}", "bledy_doby": "Doby", "bledy_przebiegi": "Przebiegi",
        "bledy_warunki_pogody": "Warunki_pogody", "najwieksze_bledy": "Najwieksze_pomylki",
        "porownanie_wariantow": "Porownanie_walidacja", "waznosc_pogody": "Waznosc_pogody",
        "pokrycie_okien": "Pokrycie_okien", "pokrycie_pogody": "Pokrycie_pogody",
        "kontrola_modelu": "Kontrola_modelu", "kontrola_energii": "Kontrola_energii",
        "pominiete_dni_zmiany_czasu": "Pominiete_doby",
        "podsumowanie_portfela": "Portfel", "backtest_bilans": "Backtest_bilans",
        "zakres_backtestu": "Zakres_backtestu", "pokrycie_backtestu": "Pokrycie_backtestu",
    }
    for key, name in sheets.items():
        if key in reports:
            sheet = add_sheet(name, reports[key])
            if "mae" in reports[key].columns and len(reports[key]):
                col = get_column_letter(reports[key].columns.get_loc("mae") + 1)
                sheet.conditional_formatting.add(f"{col}2:{col}{sheet.max_row}",
                    ColorScaleRule(start_type="min", start_color="E2F0D9", end_type="max", end_color="F8CBAD"))
    ranking = reports["bledy_oddzialy"].loc[lambda x: x["kierunek"].eq("BILANS"),
                                         ["oddzial_code", "mae_ml_wspolne", "mae_profil_wspolne"]]
    charts = add_sheet("Wykresy", ranking)
    if len(ranking):
        chart = BarChart()
        chart.title = "MAE bilansu: ML i profil na wspólnych wierszach testu"
        chart.y_axis.title = "Jednostka źródłowa energii"
        chart.add_data(Reference(charts, min_col=2, max_col=3, min_row=1, max_row=len(ranking) + 1), titles_from_data=True)
        chart.set_categories(Reference(charts, min_col=1, min_row=2, max_row=len(ranking) + 1))
        chart.width, chart.height = 25, 12
        charts.add_chart(chart, "E2")
    hourly = reports["bledy_godziny"].loc[lambda x: x["kierunek"].eq("BILANS")]
    hourly = hourly.pivot(index="godzina_lokalna", columns="oddzial_code", values="mae").reset_index()
    hourly_sheet = add_sheet("Wykres_godzinowy", hourly)
    if len(hourly) and len(hourly.columns) > 1:
        chart = LineChart()
        chart.title = "MAE bilansu według godziny lokalnej i oddziału"
        chart.x_axis.title, chart.y_axis.title = "Godzina lokalna", "MAE"
        chart.add_data(Reference(hourly_sheet, min_col=2, max_col=len(hourly.columns), min_row=1, max_row=len(hourly) + 1), titles_from_data=True)
        chart.set_categories(Reference(hourly_sheet, min_col=1, min_row=2, max_row=len(hourly) + 1))
        chart.width, chart.height = 25, 12
        hourly_sheet.add_chart(chart, "L2")
    if "przebieg_ciagly" in reports:
        for number, (branch, frame) in enumerate(reports["przebieg_ciagly"].groupby("oddzial_code"), 1):
            columns = [f"{channel}_{value}" for channel in ["bilans", "pobranie", "oddanie"]
                       for value in ["wartosc_rzeczywista", "wartosc_przewidywana", "wartosc_bazowa"]]
            timeline = frame.sort_values("target_time_utc")[["target_time_utc", *columns]].head(5000).copy()
            timeline["target_time_utc"] = timeline.target_time_utc.dt.tz_convert(bundle["settings"]["timezone"]).astype(str)
            sheet = add_sheet(f"Przebieg_{number}_{branch}", timeline)
            for channel_number, channel in enumerate(["bilans", "pobranie", "oddanie"]):
                chart = LineChart()
                chart.title = f"{branch}: {channel} — realizacja, ML, profil"
                chart.y_axis.title = "Jednostka źródłowa energii"
                chart.x_axis.title = "Czas lokalny; najnowszy origin dla celu"
                start_column = 2 + channel_number * 3
                chart.add_data(Reference(sheet, min_col=start_column, max_col=start_column + 2,
                                         min_row=1, max_row=len(timeline) + 1), titles_from_data=True)
                chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=len(timeline) + 1))
                chart.width, chart.height = 28, 13
                chart.display_blanks = "gap"
                sheet.add_chart(chart, f"L{2 + channel_number * 27}")
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp.xlsx")
    workbook.save(temporary)
    temporary.replace(path)


def print_training_summary(reports: dict[str, pd.DataFrame]):
    summary = reports["bledy_oddzialy"]
    balance = summary.loc[summary["kierunek"].eq("BILANS")].sort_values("mae", ascending=False)
    print("\nTEST — BILANS: błędy w jednostce źródłowej; dodatni bias oznacza zawyżanie.")
    print(balance[["oddzial_code", "n", "mae", "rmse", "bias", "wape", "ocena_vs_profil"]].to_string(index=False))
    worse = summary.loc[summary["ocena_vs_profil"].eq("ML_GORSZY_OD_PROFILU"), ["oddzial_code", "kierunek"]]
    if not worse.empty:
        print("UWAGA: ML przegrał z profilem historycznym w poniższych zakresach. Sprawdź dane i ocenę na nowym odłożonym okresie:")
        print(worse.to_string(index=False))
    print("Szczegóły: raport_treningu.xlsx; największe pomyłki: najwieksze_bledy.csv; wpływ pogody: waznosc_pogody.csv.")
