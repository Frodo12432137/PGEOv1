from __future__ import annotations

from contextlib import redirect_stdout
from dataclasses import asdict, replace
from io import StringIO
import json
from pathlib import Path
import sys
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from threadpoolctl import threadpool_limits

from mdd_forecasting.backtest import (dst_target_mask, run_backtest, run_signature, score_origin,
                                      settings_from_options, time_scope)
from mdd_forecasting.backtest_runtime import exclusive_run, supervise
from mdd_forecasting.branch_data import load_branch_energy
from mdd_forecasting.branch_model import TrainingSettings, load_bundle, predict_branches, train_branch_models
from mdd_forecasting.config import PipelineConfig
from tests.test_branch_forecasting import demo_data


class TestBacktest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        raw, energy, raw_weather, weather, _ = demo_data(45)
        cls.raw = raw.loc[raw.oddzial_code.eq("BIA")].copy()
        cls.energy = energy.loc[energy.oddzial_code.eq("BIA")].copy()
        cls.raw_weather = raw_weather.loc[raw_weather.punkt.eq("Białystok")].copy()
        cls.weather = weather.loc[weather.punkt.eq(cls.energy.punkt.iloc[0])].copy()
        cls.settings = TrainingSettings(validation_days=4, test_days=4, label_delay_hours=24,
                                        min_train_rows=30, min_evaluation_targets=12, max_train_rows=20_000,
                                        max_iter=4, importance_rows=0, threads=1,
                                        weather_variants=("kalendarz", "temperatura"))
        cls.train = cls.energy.loc[cls.energy.target_time_utc.lt(pd.Timestamp("2025-02-06T00:00Z"))].copy()
        with threadpool_limits(limits=1), redirect_stdout(StringIO()):
            cls.bundle, cls.internal = train_branch_models(cls.train, cls.weather, cls.settings)
        cls.bundle["input_schema"] = {"exclude_dst_days": True}
        cls.origin = pd.Timestamp("2025-02-10T00:00Z")

    def configuration(self, directory):
        energy = directory / "energia.csv"
        weather = directory / "pogoda.csv"
        self.raw.to_csv(energy, sep=";", index=False, encoding="utf-8-sig")
        self.raw_weather.to_csv(weather, sep=";", index=False, encoding="utf-8-sig")
        return {"energy": str(energy), "weather_file": str(weather),
                "mapping": str(PipelineConfig().mapping_path), "run_dir": str(directory / "run"),
                "cache_root": str(directory / "cache"), "train_from": "2025-01-01", "train_to": "2025-02-05",
                "backtest_from": "2025-02-10", "backtest_to": "2025-02-10", "parameters": asdict(self.settings),
                "weather_options": {}, "runtime": {"resume": True, "report_every_origins": 10, "max_html_origins": 20}}

    def test_dates_are_local_inclusive_with_embargo_and_strict_format(self):
        config = {"train_from": "2025-01-01", "train_to": "2025-02-05",
                  "backtest_from": "2025-02-10", "backtest_to": "2025-02-10"}
        start, end, origins = time_scope(config, self.settings)
        self.assertEqual(str(start), "2024-12-31 23:00:00+00:00")
        self.assertEqual(str(end), "2025-02-05 23:00:00+00:00")
        self.assertEqual(len(origins), 2)
        with self.assertRaisesRegex(ValueError, "embargo"):
            time_scope({**config, "train_to": "2025-02-10"}, self.settings)
        with self.assertRaises(ValueError):
            time_scope({**config, "train_from": "01.01.2025"}, self.settings)

    def test_production_predictions_identical_and_test_actuals_cannot_change_them(self):
        with threadpool_limits(limits=1):
            production, _ = predict_branches(self.bundle, self.weather, self.origin)
            balance, details = score_origin(self.bundle, self.weather, self.energy, self.origin)
            changed = self.energy.copy()
            changed.loc[changed.target_time_utc.gt(self.origin), "wartosc_rzeczywista"] += 1_000_000
            alternative, _ = score_origin(self.bundle, self.weather, changed, self.origin)
        self.assertEqual(len(balance), 72)
        for column in ("pobranie", "oddanie", "bilans"):
            np.testing.assert_array_equal(production[column], balance[column + "_wartosc_przewidywana"])
            np.testing.assert_array_equal(balance[column + "_wartosc_przewidywana"], alternative[column + "_wartosc_przewidywana"])
        self.assertFalse(balance.bilans_wartosc_rzeczywista.equals(alternative.bilans_wartosc_rzeczywista))
        matched = details.loc[details.pogoda_dopasowana]
        self.assertTrue(matched.weather_issue_time_utc.le(self.origin).all())

    def test_future_weather_issues_never_affect_backtest(self):
        modified = self.weather.copy()
        modified.loc[modified.weather_issue_time_utc.gt(self.origin), "temperatura"] = 999999
        with threadpool_limits(limits=1):
            first, _ = score_origin(self.bundle, self.weather, self.energy, self.origin)
            second, _ = score_origin(self.bundle, modified, self.energy, self.origin)
        np.testing.assert_array_equal(first.bilans_wartosc_przewidywana, second.bilans_wartosc_przewidywana)

    def test_missing_actual_is_not_zero_and_dst_days_are_calendar_based(self):
        missing = self.energy.loc[~(self.energy.target_time_utc.eq(self.origin + pd.Timedelta(hours=1))
                                   & self.energy.kierunek_energii_norm.eq("ODDANIE"))]
        with threadpool_limits(limits=1):
            balance, details = score_origin(self.bundle, self.weather, missing, self.origin)
        self.assertEqual(balance.iloc[0].status_oceny, "BRAK_REALIZACJI")
        self.assertTrue(pd.isna(balance.iloc[0].bilans_wartosc_rzeczywista))
        self.assertTrue(pd.notna(balance.iloc[0].bilans_wartosc_przewidywana))
        times = pd.Series(pd.to_datetime(["2024-03-31T01:00Z", "2024-10-27T01:00Z", "2024-04-01T10:00Z"]))
        self.assertEqual(dst_target_mask(times, "Europe/Warsaw").tolist(), [True, True, False])

    def test_completed_tuning_trials_resume_without_changing_model(self):
        from mdd_forecasting.branch_model import _new_model
        cache = {}
        with threadpool_limits(limits=1), redirect_stdout(StringIO()):
            with patch("mdd_forecasting.branch_model._new_model", wraps=_new_model) as factory:
                first, _ = train_branch_models(self.train, self.weather, self.settings, trial_cache=cache)
                first_fits = factory.call_count
            with patch("mdd_forecasting.branch_model._new_model", wraps=_new_model) as factory:
                second, _ = train_branch_models(self.train, self.weather, self.settings, trial_cache=cache)
                second_fits = factory.call_count
            a, _ = predict_branches(first, self.weather, self.origin)
            b, _ = predict_branches(second, self.weather, self.origin)
        self.assertEqual(len(cache), 4)
        self.assertEqual(first_fits, 8)
        self.assertEqual(second_fits, 4)  # Obowiązkowe refity wewnętrznego testu i końcowego modelu.
        np.testing.assert_array_equal(a.bilans, b.bilans)

    def test_full_worker_reports_reusable_model_and_no_external_labels_in_fit(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            configuration = self.configuration(root)
            actual_train = train_branch_models

            def check_training(energy, weather, settings, **kwargs):
                self.assertTrue(energy.target_time_utc.lt(pd.Timestamp("2025-02-05T23:00Z")).all())
                self.assertEqual(settings, self.settings)
                return actual_train(energy, weather, settings, **kwargs)

            with patch("mdd_forecasting.backtest.train_branch_models", side_effect=check_training) as fitted, redirect_stdout(StringIO()):
                result = run_backtest(configuration)
            self.assertEqual(result, 0)
            self.assertEqual(fitted.call_count, 1)
            directory = root / "run"
            status = json.loads((directory / "status_backtestu.json").read_text(encoding="utf-8"))
            self.assertTrue(status["kompletny"])
            self.assertEqual(status["gotowe_prognozy"], 2)
            balance = pd.read_csv(directory / "backtest_bilans.csv", sep=";", decimal=",")
            self.assertEqual(len(balance), 144)
            self.assertTrue(balance.status_oceny.eq("OCENIONO").all())
            workbook = load_workbook(directory / "raport_backtestu.xlsx")
            self.assertIn("Zakres_backtestu", workbook.sheetnames)
            self.assertIn("Pokrycie_backtestu", workbook.sheetnames)
            curves = workbook["Przebieg_1_BIA"]
            self.assertEqual(len(curves._charts), 3)
            self.assertIn("zamrożony", workbook["Czytaj_najpierw"]["B3"].value)
            workbook.close()
            html = (directory / "raport_backtestu.html").read_text(encoding="utf-8")
            self.assertIn('application/json', html)
            self.assertNotIn('src="https://', html)
            self.assertNotIn("__PAYLOAD__", html)
            model = load_bundle(directory / "trening")
            self.assertLess(pd.Timestamp(model["training_target_max_utc"]), self.origin)
            second = {**configuration, "run_dir": str(root / "run2")}
            with patch("mdd_forecasting.backtest.train_branch_models", side_effect=AssertionError("Ponowny fit")), redirect_stdout(StringIO()):
                self.assertEqual(run_backtest(second), 0)
            pd.testing.assert_frame_equal(balance, pd.read_csv(root / "run2" / "backtest_bilans.csv", sep=";", decimal=","))
            old_signature = run_signature(configuration)
            with Path(configuration["energy"]).open("a", encoding="utf-8") as stream:
                stream.write("\n")
            self.assertNotEqual(old_signature, run_signature(configuration))

    def test_production_launcher_options_remain_identical(self):
        from uruchom_trening_bilansow import PARAMETRY
        from uruchom_backtest_bilansow import PARAMETRY as BACKTEST_PARAMETERS
        self.assertEqual(PARAMETRY, BACKTEST_PARAMETERS)
        self.assertEqual(settings_from_options(PARAMETRY).max_iter, 300)
        self.assertTrue(settings_from_options(PARAMETRY).tune)

    def test_end_date_filters_before_bad_hour_validation(self):
        raw = self.raw.head(2).copy()
        bad = raw.iloc[[0]].copy()
        bad["doba_handlowa"] = pd.Timestamp("2026-01-01")
        bad["godzina_handlowa"] = 999
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "energy.csv"
            pd.concat([raw, bad]).to_csv(path, index=False, sep=";")
            result, quality, _ = load_branch_energy(path, PipelineConfig(), end_date="2025-12-31")
        self.assertFalse(result.empty)
        self.assertEqual(quality.loc[quality.kontrola.eq("po_dacie_koncowej"), "liczba"].iloc[0], 1)

    def test_time_budget_stops_child_and_marks_result_incomplete(self):
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            with redirect_stdout(StringIO()):
                code = supervise([sys.executable, "-c", "import time; time.sleep(30)"], root, .1)
            self.assertEqual(code, 124)
            status = json.loads((root / "status_backtestu.json").read_text(encoding="utf-8"))
            self.assertEqual(status["etap"], "LIMIT_CZASU")
            self.assertFalse(status["kompletny"])

    def test_output_lock_prevents_concurrent_cache_writes(self):
        with TemporaryDirectory() as tmp:
            with exclusive_run(Path(tmp)):
                with self.assertRaisesRegex(RuntimeError, "Inny backtest"):
                    with exclusive_run(Path(tmp)):
                        pass


if __name__ == "__main__":
    unittest.main()
