from __future__ import annotations

from contextlib import redirect_stdout
from copy import deepcopy
from dataclasses import asdict, replace
from io import StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from threadpoolctl import threadpool_limits

from mdd_forecasting.branch_data import (aggregate_branch_energy, historical_requests, join_weather_asof,
                                        load_prediction_measurements, sum_branch_measurements)
from mdd_forecasting.branch_launcher import _options
from mdd_forecasting.branch_model import balance_table, train_branch_models
from mdd_forecasting.energy_filter import (EnergyFilterSettings, apply_energy_filter, fit_energy_filter,
                                          prepare_training_energy, training_split_cutoff)
from tests.test_branch_forecasting import normalized_energy
from tests import test_backtest as fixture


def measurement_rows(days=12):
    rows = []
    for day in pd.date_range("2025-01-01", periods=days, freq="D"):
        for hour in range(1, 25):
            rows.append(["BIA", 1, "G", "Małe MDD", "", day, 4, hour, 100.])
            rows.append(["BIA", -1, "G", "Małe MDD", "", day, 4, hour, 20. if 9 <= hour <= 18 else 0.])
            rows.append(["BIA", 1, "G", "Duże MDD", "", day, 4, hour, 10000.])
    return rows


class TestEnergyFilter(unittest.TestCase):
    def selected(self, rows):
        with redirect_stdout(StringIO()):
            return aggregate_branch_energy(normalized_energy(rows), return_measurements=True,
                                           exclude_negative_energy=True)

    def test_negative_best_revision_is_missing_not_older_revision_or_partial_sum(self):
        rows = [["BIA", 1, "G", "A", "", "2025-01-01", 3, 1, 100],
                ["BIA", 1, "G", "A", "", "2025-01-01", 4, 1, -5],
                ["BIA", 1, "G", "B", "", "2025-01-01", 4, 1, 20],
                ["BIA", -1, "G", "A", "", "2025-01-01", 4, 1, 50]]
        energy, quality, measurements = self.selected(rows)
        withdrawal = energy.loc[energy.kierunek_energii_norm.eq("POBRANIE")].iloc[0]
        self.assertTrue(pd.isna(withdrawal.wartosc_rzeczywista))
        self.assertEqual(withdrawal.liczba_mdd, 2)
        self.assertEqual(withdrawal.liczba_odrzuconych_pomiarow, 1)
        invalid = measurements.loc[measurements.powod_wykluczenia_energii.eq("UJEMNY_POMIAR")].iloc[0]
        self.assertEqual(invalid.wartosc_przed_filtracja, -5)
        self.assertEqual(invalid.rodzaj_rank, 4)
        self.assertEqual(quality.loc[quality.kontrola.eq("odrzucone_ujemne_pomiary"), "liczba"].iloc[0], 1)

    def test_old_negative_revision_does_not_block_valid_best_correction_and_negative_balance_is_valid(self):
        rows = [["BIA", 1, "G", "A", "", "2025-01-01", 2, 1, -100],
                ["BIA", 1, "G", "A", "", "2025-01-01", 4, 1, 0],
                ["BIA", -1, "G", "A", "", "2025-01-01", 4, 1, 50]]
        with redirect_stdout(StringIO()):
            energy, _ = aggregate_branch_energy(normalized_energy(rows))
        energy["forecast_origin_utc"] = energy.target_time_utc - pd.Timedelta(hours=12)
        energy["horyzont_h"] = 12
        balance = balance_table(energy, ["wartosc_rzeczywista"])
        self.assertEqual(balance.bilans_wartosc_rzeczywista.iloc[0], -50)

    def test_frozen_thresholds_are_separate_by_mdd_and_direction_and_keep_zero(self):
        _, _, measurements = self.selected(measurement_rows())
        settings = EnergyFilterSettings(energy_outlier_min_samples=4)
        before = pd.Timestamp("2025-01-10T00:00Z")
        state = fit_energy_filter(measurements, before, settings, "Europe/Warsaw")
        future = measurements.loc[measurements.target_time_utc.ge(before)].copy()
        first = future.loc[future.klient_nazwa.eq("Małe MDD") & future.kierunek_energii_norm.eq("POBRANIE")].index[0]
        second = future.loc[future.klient_nazwa.eq("Duże MDD")].index[0]
        future.loc[first, ["wartosc_rzeczywista", "wartosc_przed_filtracja"]] = 5000.
        future.loc[second, ["wartosc_rzeczywista", "wartosc_przed_filtracja"]] = 15000.
        state_copy = deepcopy(state)
        filtered, reports = apply_energy_filter(future, state, "Europe/Warsaw")
        self.assertEqual(state, state_copy)
        self.assertTrue(pd.isna(filtered.loc[first, "wartosc_rzeczywista"]))
        self.assertEqual(filtered.loc[first, "prog_gorny_energii"], 300.)
        self.assertEqual(filtered.loc[second, "wartosc_rzeczywista"], 15000.)
        self.assertTrue(filtered.loc[filtered.wartosc_przed_filtracja.eq(0), "wartosc_rzeczywista"].eq(0).all())
        self.assertEqual(len(reports["odrzucone_pomiary_energii"]), 1)
        hourly = sum_branch_measurements(filtered, "Europe/Warsaw")
        self.assertTrue(pd.isna(hourly.loc[hourly.target_time_utc.eq(filtered.loc[first, "target_time_utc"])
                                         & hourly.kierunek_energii_norm.eq("POBRANIE"), "wartosc_rzeczywista"].iloc[0]))

    def test_validation_values_cannot_move_learned_limits(self):
        _, _, measurements = self.selected(measurement_rows())
        before = pd.Timestamp("2025-01-10T00:00Z")
        settings = EnergyFilterSettings(energy_outlier_min_samples=4)
        state = fit_energy_filter(measurements, before, settings, "Europe/Warsaw")
        changed = measurements.copy()
        changed.loc[changed.target_time_utc.ge(before), ["wartosc_rzeczywista", "wartosc_przed_filtracja"]] = 1e12
        self.assertEqual(state, fit_energy_filter(changed, before, settings, "Europe/Warsaw"))

    def test_new_mdd_and_zero_only_reference_do_not_get_an_invented_threshold(self):
        _, _, measurements = self.selected(measurement_rows(2))
        before = pd.Timestamp("2025-01-02T00:00Z")
        zero = measurements.copy()
        zero[["wartosc_rzeczywista", "wartosc_przed_filtracja"]] = 0.
        state = fit_energy_filter(zero, before, EnergyFilterSettings(), "Europe/Warsaw")
        self.assertEqual(state["thresholds"], [])
        sample = measurements.tail(2).copy()
        sample["klient_nazwa"] = "NOWE_MDD"
        sample[["wartosc_rzeczywista", "wartosc_przed_filtracja"]] = 1e9
        filtered, reports = apply_energy_filter(sample, state, "Europe/Warsaw")
        self.assertTrue(filtered.wartosc_rzeczywista.eq(1e9).all())
        self.assertTrue(reports["odrzucone_pomiary_energii"].empty)

    def test_filter_can_be_disabled_and_launcher_keeps_false_options(self):
        _, _, measurements = self.selected(measurement_rows())
        state = fit_energy_filter(measurements, pd.Timestamp("2025-01-10T00:00Z"),
                                  EnergyFilterSettings(exclude_energy_outliers=False), "Europe/Warsaw")
        measurements.iloc[0, measurements.columns.get_loc("wartosc_przed_filtracja")] = 1e12
        result, reports = apply_energy_filter(measurements, state, "Europe/Warsaw")
        self.assertEqual(result.wartosc_rzeczywista.iloc[0], 1e12)
        self.assertTrue(result.status_progu_energii.eq("FILTR_ODSTAJACYCH_WYLACZONY").all())
        self.assertEqual(_options({"exclude_negative_energy": False, "exclude_energy_outliers": False}),
                         ["--no-exclude-negative-energy", "--no-exclude-energy-outliers"])

    def test_missing_values_remain_missing_without_ambiguous_nullable_masks(self):
        rows = measurement_rows()
        rows[0][-1], rows[1][-1] = None, None
        energy, _, measurements = self.selected(rows)
        state = fit_energy_filter(measurements, pd.Timestamp("2025-01-10T00:00Z"),
                                  EnergyFilterSettings(energy_outlier_min_samples=4), "Europe/Warsaw")
        filtered, reports = apply_energy_filter(measurements, state, "Europe/Warsaw")
        self.assertTrue(filtered.loc[filtered.wartosc_przed_filtracja.isna(), "wartosc_rzeczywista"].isna().all())
        self.assertTrue(reports["odrzucone_pomiary_energii"].empty)
        self.assertTrue(energy.iloc[0:2].wartosc_rzeczywista.isna().all())


class TestFilterTrainingBoundary(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        fixture.TestBacktest.setUpClass()

    def test_reference_anchor_matches_historical_requests_for_different_grids_and_weather_end_dates(self):
        energy, weather = fixture.TestBacktest.energy, fixture.TestBacktest.weather
        for horizon, step, last in [(72, 12, "2025-02-14T23:00Z"), (6, 12, "2025-02-03T11:00Z"), (24, 7, "2025-02-04T07:00Z")]:
            with self.subTest(horizon=horizon, step=step):
                settings = replace(fixture.TestBacktest.settings, horizon_hours=horizon, origin_step_hours=step)
                wx = weather.loc[weather.weather_valid_time_utc.le(pd.Timestamp(last))]
                requests = historical_requests(energy, horizon, step)
                joined = join_weather_asof(requests, wx)
                expected = joined.loc[joined.pogoda_dopasowana & joined.wartosc_rzeczywista.notna(), "target_time_utc"].max().floor("D")
                self.assertEqual(training_split_cutoff(energy, wx, settings), expected)

    def test_training_keeps_split_anchor_when_last_day_is_filtered(self):
        with redirect_stdout(StringIO()):
            _, _, measurements = aggregate_branch_energy(normalized_energy(fixture.TestBacktest.raw.values.tolist()),
                                                         return_measurements=True, exclude_negative_energy=True)
        last_day = measurements.target_time_utc.dt.floor("D").max()
        measurements.loc[measurements.target_time_utc.ge(last_day), ["wartosc_rzeczywista", "wartosc_przed_filtracja"]] = 1e9
        energy = sum_branch_measurements(measurements, "Europe/Warsaw")
        settings = replace(fixture.TestBacktest.settings, weather_variants=("kalendarz",), max_iter=2)
        cleaned, _, state, reports, cutoff = prepare_training_energy(
            energy, measurements, fixture.TestBacktest.weather, settings, EnergyFilterSettings(energy_outlier_min_samples=4))
        self.assertEqual(cutoff, last_day)
        self.assertTrue(cleaned.loc[cleaned.target_time_utc.ge(last_day), "wartosc_rzeczywista"].isna().all())
        with redirect_stdout(StringIO()), threadpool_limits(limits=1):
            _, model_reports = train_branch_models(cleaned, fixture.TestBacktest.weather, settings, split_cutoff=cutoff)
        expected = cutoff - pd.Timedelta(days=settings.test_days)
        actual = model_reports["kontrola_modelu"].loc[lambda x: x.kontrola.eq("poczatek_testu"), "szczegoly"].iloc[0]
        self.assertEqual(pd.Timestamp(actual), expected)
        self.assertLess(pd.Timestamp(state["reference_before_utc"]), expected)

    def test_dirty_backtest_runs_reports_exclusions_and_prediction_audit_uses_saved_thresholds(self):
        from mdd_forecasting.backtest import run_backtest
        from mdd_forecasting.branch_model import load_bundle
        with TemporaryDirectory() as tmp:
            root = Path(tmp)
            configuration = fixture.TestBacktest().configuration(root)
            configuration["parameters"].update(energy_outlier_min_samples=10)
            raw = fixture.TestBacktest.raw.copy()
            for date, hour, value in [("2025-01-10", 3, -10.), ("2025-01-12", 4, 1e9),
                                      ("2025-02-10", 3, -10.), ("2025-02-10", 4, 1e9)]:
                mask = raw.doba_handlowa.eq(pd.Timestamp(date)) & raw.godzina_handlowa.eq(hour) & raw.kierunek_code.eq(1)
                self.assertEqual(int(mask.sum()), 1)
                raw.loc[mask, "wartosc_rzeczywista"] = value
            source = Path(configuration["energy"])
            raw.to_csv(source, sep=";", index=False)
            original_bytes = source.read_bytes()
            with redirect_stdout(StringIO()):
                code = run_backtest(configuration)
            self.assertEqual(code, 2)
            run = root / "run"
            balance = pd.read_csv(run / "backtest_bilans.csv", sep=";", decimal=",")
            self.assertEqual(len(balance), 144)
            self.assertEqual(int(balance.status_oceny.eq("POMIAR_WYKLUCZONY_FILTREM").sum()), 2)
            self.assertTrue(balance.bilans_wartosc_przewidywana.notna().all())
            discarded = pd.read_csv(run / "odrzucone_pomiary_energii.csv", sep=";", decimal=",")
            self.assertEqual(len(discarded), 4)
            self.assertEqual(set(discarded.powod_wykluczenia_energii), {"UJEMNY_POMIAR", "ODSTAJACY_POMIAR_POWYZEJ_PROGU"})
            workbook = load_workbook(run / "raport_backtestu.xlsx", read_only=True)
            self.assertIn("Odrzucone_pomiary", workbook.sheetnames)
            self.assertIn("Godziny_wykluczone_filtrem", workbook.sheetnames)
            workbook.close()
            bundle = load_bundle(run / "trening")
            state = deepcopy(bundle["energy_filter"])
            with redirect_stdout(StringIO()):
                _, _, reports = load_prediction_measurements(source, bundle, return_filter_reports=True)
            self.assertEqual(bundle["energy_filter"], state)
            self.assertEqual(len(reports["odrzucone_pomiary_energii"]), 4)
            self.assertEqual(source.read_bytes(), original_bytes)
            self.assertTrue(json.loads((run / "status_backtestu.json").read_text())["kompletny"])
            from mdd_forecasting.branch_cli import train_main
            cli_directory = root / "zwykly_trening"
            with redirect_stdout(StringIO()):
                self.assertEqual(train_main([
                    "--energy", str(source), "--weather", configuration["weather_file"],
                    "--artifacts-dir", str(cli_directory), "--start-date", "2025-01-01",
                    "--validation-days", "4", "--test-days", "4", "--label-delay-hours", "24",
                    "--max-iter", "2", "--importance-rows", "0", "--threads", "1",
                    "--weather-variants", "kalendarz",
                ]), 0)
            cli_bundle = load_bundle(cli_directory)
            self.assertTrue(cli_bundle["input_schema"]["exclude_negative_energy"])
            rejected = pd.read_csv(cli_directory / "odrzucone_pomiary_energii.csv", sep=";", decimal=",")
            self.assertEqual(len(rejected), 4)


if __name__ == "__main__":
    unittest.main()
