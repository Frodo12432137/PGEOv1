from __future__ import annotations

from dataclasses import replace
from importlib.util import find_spec
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import hashlib
import os
import unittest

import numpy as np
import pandas as pd
from threadpoolctl import threadpool_limits

from mdd_forecasting.branch_data import aggregate_branch_energy
from mdd_forecasting.branch_diagnostics import error_statistics, write_training_workbook
from mdd_forecasting.branch_launcher import configured_launch
from mdd_forecasting.branch_model import TrainingSettings, train_branch_models, predict_branches, save_bundle, load_bundle, metrics_table
from tests.test_branch_forecasting import demo_data, normalized_energy


class TestAdditionalEnergyChecks(unittest.TestCase):
    def test_reactive_energy_grammatical_forms_are_rejected(self):
        for label, sign in [("energia bierna pobrana", 1), ("energia bierna oddana", -1), ("bierne pobranie", 1)]:
            with self.subTest(label=label), self.assertRaisesRegex(ValueError, "biernej"):
                aggregate_branch_energy(normalized_energy([["BIA", sign, "G", "K1", label, "2025-01-01", 4, 1, 10]]))

    def test_absent_record_within_day_invalidates_partial_total(self):
        rows = [["BIA", 1, "G", client, "czynne pobranie", "2025-01-01", 4, hour, value]
                for hour in [1, 2, 3] for client, value in [("K1", 10), ("K2", 90)]
                if not (hour == 2 and client == "K2")]
        hourly, quality = aggregate_branch_energy(normalized_energy(rows))
        middle = hourly.loc[hourly.godzina_handlowa.eq(2)].iloc[0]
        self.assertTrue(pd.isna(middle.wartosc_rzeczywista))
        self.assertEqual(middle.brakujace_rekordy_mdd, 1)
        self.assertEqual(quality.loc[quality.kontrola.eq("luki_godzinowe_mdd_w_dobie"), "liczba"].iloc[0], 1)

    def test_whole_missing_hour_is_retained_as_missing_and_autumn_hours_are_distinct(self):
        rows = [["BIA", 1, "G", "K1", "czynne pobranie", "2025-10-26", 4, hour, 10] for hour in [2, 4]]
        hourly, _ = aggregate_branch_energy(normalized_energy(rows))
        missing = hourly.loc[hourly.wartosc_rzeczywista.isna()]
        self.assertEqual(set(missing.godzina_handlowa), {3, 25})
        self.assertEqual(missing.target_time_utc.nunique(), 2)

    def test_no_contract_inference_outside_observed_span(self):
        rows = [["BIA", 1, "G", "K1", "czynne pobranie", "2025-01-01", 4, hour, 0] for hour in [10, 11]]
        hourly, _ = aggregate_branch_energy(normalized_energy(rows))
        self.assertEqual(set(hourly.godzina_handlowa), {10, 11})
        self.assertTrue(hourly.wartosc_rzeczywista.eq(0).all())

    def test_infinite_energy_is_rejected(self):
        row = ["BIA", 1, "G", "K1", "czynne pobranie", "2025-01-01", 4, 1, np.inf]
        with self.assertRaisesRegex(ValueError, "skończoną"):
            aggregate_branch_energy(normalized_energy([row]))


class TestTrainingTools(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw, cls.energy, cls.weather_raw, cls.weather, cls.origin = demo_data(days=25)
        cls.settings = TrainingSettings(validation_days=4, test_days=4, label_delay_hours=24,
                                       max_iter=5, min_train_rows=40, importance_rows=12,
                                       weather_variants=("temperatura_slonce",), threads=2)
        with threadpool_limits(limits=2):
            cls.bundle, cls.reports = train_branch_models(cls.energy, cls.weather, cls.settings)

    def test_missing_validation_and_test_branch_prevents_model(self):
        incomplete = self.energy.loc[~(self.energy.oddzial_code.eq("BIA")
                                      & self.energy.target_time_utc.ge(self.origin - pd.Timedelta(days=10)))]
        with self.assertRaisesRegex(ValueError, "pokrycia.*BIA"):
            train_branch_models(incomplete, self.weather, self.settings)

    def test_partial_weather_is_marked_and_missing_weather_remains_missing(self):
        partial = self.weather.copy()
        partial[["promieniowanie_calkowite", "promieniowanie_bezposrednie"]] = np.nan
        with threadpool_limits(limits=2):
            balance, details = predict_branches(self.bundle, partial, self.origin)
        self.assertEqual(len(balance), 576)
        self.assertTrue(balance.status.eq("PROGNOZA_NIEPELNA_POGODA").all())
        self.assertTrue(balance.bilans.notna().all())
        self.assertTrue(details.brakujace_cechy_pogody.str.contains("promieniowanie_calkowite").all())
        empty = partial.copy()
        empty["temperatura"] = np.nan
        with threadpool_limits(limits=2):
            balance, details = predict_branches(self.bundle, empty, self.origin)
        self.assertTrue(balance.bilans.isna().all())
        self.assertTrue(details.status.eq("BRAK_POGODY_PRZED_URUCHOMIENIEM").all())

    def test_statistics_use_paired_baseline_and_handle_zero_targets(self):
        group = pd.DataFrame({"realizacja": [0., 10., 20.], "prognoza": [0., 12., 16.],
                              "profil": [0., 15., np.nan],
                              "target_time_utc": pd.date_range("2025-01-01", periods=3, freq="h", tz="UTC")})
        metrics = error_statistics(group)
        self.assertAlmostEqual(metrics["mae"], 2)
        self.assertAlmostEqual(metrics["bias"], -2 / 3)
        self.assertAlmostEqual(metrics["mae_ml_wspolne"], 1)
        self.assertAlmostEqual(metrics["mae_profil_wspolne"], 2.5)
        self.assertAlmostEqual(metrics["poprawa_mae_vs_profil"], .6)
        self.assertTrue(np.isfinite(metrics["smape"]))
        zero = group.assign(realizacja=0., prognoza=0., profil=0.)
        zero_metrics = error_statistics(zero)
        self.assertEqual(zero_metrics["smape"], 0)
        self.assertTrue(np.isnan(zero_metrics["wape"]))
        self.assertTrue(np.isnan(zero_metrics["r2"]))

    def test_short_horizon_metrics_do_not_duplicate_total_scope(self):
        frame = pd.DataFrame({"oddzial_code": ["BIA"] * 24, "horyzont_h": range(1, 25),
                              "actual": np.ones(24), "prediction": np.ones(24)})
        metrics = metrics_table(frame, "actual", "prediction", "ML")
        self.assertEqual(metrics.horyzont.tolist(), ["1-24"])
        self.assertEqual(metrics.iloc[0]["n"], 24)

    def test_nullable_missing_realizations_are_reported_without_crashing(self):
        group = pd.DataFrame({"oddzial_code": ["BIA", "BIA"], "horyzont_h": [1, 2],
                              "realizacja": pd.Series([pd.NA, 10], dtype="Float64"),
                              "prognoza": pd.Series([5, 12], dtype="Float64"),
                              "profil": pd.Series([5, pd.NA], dtype="Float64"),
                              "target_time_utc": pd.date_range("2025-01-01", periods=2, freq="h", tz="UTC")})
        stats = error_statistics(group)
        self.assertEqual(stats["n"], 1)
        self.assertEqual(stats["n_wspolne_z_profilem"], 0)
        self.assertEqual(stats["pokrycie"], .5)
        self.assertEqual(stats["mae"], 2)
        metrics = metrics_table(group, "realizacja", "prognoza", "ML")
        self.assertEqual(metrics.iloc[0]["n"], 1)

    def test_workbook_contains_diagnostics_charts_and_safe_text(self):
        from openpyxl import load_workbook
        reports = dict(self.reports)
        reports["kontrola_energii"] = pd.DataFrame({"opis": ["=1+1"]})
        with TemporaryDirectory() as tmp:
            path = Path(tmp) / "raport.xlsx"
            write_training_workbook(self.bundle, reports, path)
            workbook = load_workbook(path, data_only=False)
            self.assertTrue({"Podsumowanie", "Najwieksze_pomylki", "Waznosc_pogody", "Warunki_pogody", "Pokrycie_okien"}.issubset(workbook.sheetnames))
            self.assertEqual(workbook["Kontrola_energii"]["A2"].data_type, "s")
            self.assertEqual(workbook["Kontrola_energii"]["A2"].value, "=1+1")
            self.assertEqual(len(workbook["Wykresy"]._charts), 1)
            self.assertEqual(len(workbook["Wykres_godzinowy"]._charts), 1)
            self.assertEqual(workbook["Podsumowanie"].max_row, 25)
            workbook.close()

    def test_configured_launch_works_outside_project_and_predicts_without_fit(self):
        with TemporaryDirectory() as tmp, TemporaryDirectory() as other:
            directory = Path(tmp)
            energy = directory / "energia.xlsx"
            weather = directory / "pogoda.csv"
            self.raw.to_excel(energy, index=False, sheet_name="dane_po_transpozycji")
            self.weather_raw.to_csv(weather, index=False)
            options = {"max_iter": 3, "validation_days": 4, "test_days": 4, "label_delay_hours": 24,
                       "weather_variants": ["temperatura_slonce"], "importance_rows": 0, "threads": 2}
            previous_cwd = Path.cwd()
            try:
                os.chdir(other)
                with patch("mdd_forecasting.branch_launcher.ROOT", directory), patch.dict("sys.modules", {"tkinter": None}):
                    self.assertEqual(configured_launch(True, "energia.xlsx", "model", "pogoda.csv", options=options), 0)
                    self.assertTrue((directory / "model" / "raport_treningu.xlsx").is_file())
                    path = directory / "model" / "modele_oddzialow.joblib"
                    checksum = hashlib.sha256(path.read_bytes()).hexdigest()
                    from sklearn.ensemble import HistGradientBoostingRegressor
                    with patch.object(HistGradientBoostingRegressor, "fit", side_effect=AssertionError("fit podczas predykcji")):
                        self.assertEqual(configured_launch(False, artifacts_dir="model", weather_file="pogoda.csv",
                                                          as_of=self.origin.isoformat(), output_dir="wynik"), 0)
                    output = pd.read_csv(directory / "wynik" / "prognoza_bilansu_72h.csv", sep=";", decimal=",")
                    self.assertEqual(len(output), 576)
                    self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), checksum)
                    partial = self.weather_raw.drop(columns=["calkowitePromieniowanieSloneczneGodzinowe", "bezposredniePromieniowanieSloneczneGodzinowe"])
                    partial.to_csv(weather, index=False)
                    self.assertEqual(configured_launch(False, artifacts_dir="model", weather_file="pogoda.csv",
                                                      as_of=self.origin.isoformat(), output_dir="niepelny"), 2)
            finally:
                os.chdir(previous_cwd)

    @unittest.skipUnless(find_spec("catboost") is not None, "Opcjonalny CatBoost nie jest zainstalowany")
    def test_real_catboost_comparison_serialization_and_prediction(self):
        settings = replace(self.settings, model_backend="auto", tune=True, importance_rows=0,
                           max_iter=3, catboost_depth=3, weather_variants=("temperatura_slonce",))
        with threadpool_limits(limits=2):
            bundle, reports = train_branch_models(self.energy, self.weather, settings)
        self.assertEqual(set(reports["porownanie_wariantow"]["backend"]), {"catboost", "hist_gradient_boosting"})
        # Wymuszona rodzina sprawdza też zapis i odczyt CatBoost, niezależnie od zwycięzcy auto.
        with threadpool_limits(limits=2):
            bundle, _ = train_branch_models(self.energy, self.weather, replace(settings, model_backend="catboost", tune=False))
        with TemporaryDirectory() as tmp:
            save_bundle(bundle, tmp)
            restored = load_bundle(tmp)
            self.assertIn("catboost", restored["versions"])
            from catboost import CatBoostRegressor
            with patch.object(CatBoostRegressor, "fit", side_effect=AssertionError("fit podczas predykcji")), threadpool_limits(limits=2):
                prediction, _ = predict_branches(restored, self.weather, self.origin)
            self.assertEqual(len(prediction), 576)
            self.assertTrue(prediction.bilans.notna().all())


if __name__ == "__main__":
    unittest.main()
