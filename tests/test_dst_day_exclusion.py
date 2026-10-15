from contextlib import redirect_stdout
from io import StringIO
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch
import unittest

import numpy as np
import pandas as pd

from mdd_forecasting.branch_cli import train_main
from mdd_forecasting.branch_data import (
    EnergyTimeError, historical_requests, load_branch_energy, load_prediction_measurements,
)
from mdd_forecasting.branch_model import load_bundle
from mdd_forecasting.config import ENERGY_COLUMNS, PipelineConfig
from mdd_forecasting.portfolio import analyze_portfolio
from tests.test_portfolio import measurements as portfolio_measurements


def energy_row(date, hour=1, *, branch="BIA", client="MDD", direction=1, revision=4, value=10):
    return [branch, direction, "G", client, "opis nie wpływa na kierunek", date, revision, hour, value]


class TestDstDayExclusion(unittest.TestCase):
    def load_rows(self, rows, **kwargs):
        with TemporaryDirectory() as tmp, redirect_stdout(StringIO()):
            source = Path(tmp) / "energia.csv"
            pd.DataFrame(rows, columns=ENERGY_COLUMNS).to_csv(source, index=False)
            return load_branch_energy(source, PipelineConfig(), return_measurements=True,
                                      exclude_dst_days=True, return_excluded_days=True, **kwargs)

    def test_calendar_excludes_spring_and_autumn_for_every_year(self):
        dates = ["2023-03-26", "2023-10-29", "2024-03-31", "2024-10-27", "2025-03-30", "2025-10-26"]
        rows = [energy_row(date, hour) for date in dates for hour in [1, 3, 25]]
        rows += [energy_row("2023-03-25", value=0), energy_row("2025-10-27", direction=-1, value=7)]
        hourly, quality, _, selected, excluded = self.load_rows(rows)
        self.assertEqual(excluded.doba_handlowa.dt.strftime("%Y-%m-%d").tolist(), dates)
        self.assertEqual(excluded.godziny_doby.tolist(), [23, 25, 23, 25, 23, 25])
        self.assertEqual(excluded.rekordy_wykluczone.tolist(), [3] * 6)
        self.assertEqual(selected.wartosc_rzeczywista.tolist(), [0, 7])
        self.assertEqual(len(hourly), 2)
        self.assertEqual(quality.loc[quality.kontrola.eq("rekordy_wykluczone_zmiana_czasu"), "liczba"].iloc[0], 18)

    def test_calendar_rule_does_not_depend_on_h25_or_number_of_records(self):
        rows = [energy_row("2025-10-26", 1)]
        rows += [energy_row("2025-10-27", hour) for hour in range(1, 24)]
        hourly, _, _, selected, excluded = self.load_rows(rows)
        self.assertEqual(excluded.rekordy_wykluczone.tolist(), [1])
        self.assertEqual(len(hourly), 23)
        self.assertEqual(selected.godzina_handlowa.tolist(), list(range(1, 24)))

    def test_full_day_is_removed_before_mapping_direction_revision_and_value_validation(self):
        rows = [energy_row("2025-03-30", 3, branch="NIEZNANY", client=None, direction=2, revision=9, value=-8),
                energy_row("2025-03-31", 1)]
        hourly, _, _, selected, excluded = self.load_rows(rows)
        self.assertEqual(selected.source_row.tolist(), [3])
        self.assertEqual(hourly.wartosc_rzeczywista.tolist(), [10])
        self.assertEqual(excluded.rekordy_wykluczone.tolist(), [1])

    def test_ordinary_days_and_unknown_dates_still_undergo_validation(self):
        cases = [
            (energy_row("zła data"), EnergyTimeError, "nierozpoznana"),
            (energy_row("2025-03-31", 25), EnergyTimeError, "godzina 25"),
            (energy_row("2025-03-31", direction=2), ValueError, "kolumnie B"),
            (energy_row("2025-03-31", value=-1), ValueError, "ujemną wartość"),
        ]
        for row, error, message in cases:
            with self.subTest(row=row), self.assertRaisesRegex(error, message):
                self.load_rows([row])

    def test_start_date_is_applied_before_dst_report_and_filter(self):
        rows = [energy_row("2024-03-31", 3), energy_row("2024-10-27", 25), energy_row("2024-10-28")]
        _, quality, _, _, excluded = self.load_rows(rows, start_date="2024-07-01")
        self.assertEqual(excluded.doba_handlowa.tolist(), [pd.Timestamp("2024-10-27")])
        self.assertEqual(quality.loc[quality.kontrola.eq("przed_data_startowa"), "liczba"].iloc[0], 1)

    def test_gaps_are_not_filled_and_actual_utc_spacing_is_preserved(self):
        rows = [energy_row("2025-03-29", value=8), energy_row("2025-03-30", 3), energy_row("2025-03-31", value=9)]
        hourly, _, _, _, excluded = self.load_rows(rows)
        self.assertEqual(len(hourly), 2)
        self.assertEqual(hourly.wartosc_rzeczywista.tolist(), [8, 9])
        self.assertEqual(hourly.target_time_utc.diff().dropna().iloc[0], pd.Timedelta(hours=47))
        requests = historical_requests(hourly, horizon=72, step_hours=1)
        self.assertFalse(requests.doba_handlowa.isin(excluded.doba_handlowa).any())
        self.assertEqual(set(requests.wartosc_rzeczywista), {8, 9})

    def test_ordinary_day_corrections_and_zeros_are_unchanged(self):
        rows = [energy_row("2025-03-29", revision=1, value=100),
                energy_row("2025-03-29", revision=4, value=0),
                energy_row("2025-03-29", revision=4, value=0),
                energy_row("2025-03-30", 3)]
        hourly, _, _, selected, _ = self.load_rows(rows)
        self.assertEqual(hourly.wartosc_rzeczywista.tolist(), [0])
        self.assertEqual(len(selected), 1)
        self.assertEqual(selected.rodzaj_rank.tolist(), [4])

    def test_prediction_audit_uses_saved_rule_and_old_bundles_keep_legacy_behavior(self):
        bundle = {"branches": pd.DataFrame({"oddzial_code": ["BIA"], "punkt": ["Białystok"]}),
                  "settings": {"timezone": "Europe/Warsaw"}, "input_schema": {"exclude_dst_days": True}}
        rows = [energy_row("2025-03-29"), energy_row("2025-03-30", 3), energy_row("2025-03-31")]
        with TemporaryDirectory() as tmp, redirect_stdout(StringIO()):
            source = Path(tmp) / "energia.csv"
            pd.DataFrame(rows, columns=ENERGY_COLUMNS).to_csv(source, index=False)
            selected, excluded = load_prediction_measurements(source, bundle, return_excluded_days=True)
            self.assertEqual(len(selected), 2)
            self.assertEqual(excluded.rekordy_wykluczone.tolist(), [1])
            bundle.pop("input_schema")
            with self.assertRaises(EnergyTimeError):
                load_prediction_measurements(source, bundle)

    def test_portfolio_does_not_count_intentionally_omitted_day_as_missing_measurements(self):
        rows = [["BIA", "POBRANIE", "A", date, 0] for date in ["2025-03-29", "2025-03-31"]]
        data = portfolio_measurements(rows)
        baseline = analyze_portfolio(data)
        self.assertEqual(baseline["przerwy_pomiarow_mdd"].dni_przerwy.tolist(), [1])
        reports = analyze_portfolio(data, ignored_dates=["2025-03-30"])
        self.assertTrue(reports["przerwy_pomiarow_mdd"].empty)
        self.assertEqual(reports["zmiany_portfela_dzienne"].status.iloc[-1], "BEZ_ZMIANY")
        self.assertEqual(reports["portfel_mdd"].liczba_przerw.tolist(), [0])

    def test_portfolio_keeps_real_gaps_and_membership_changes_across_omitted_day(self):
        data = portfolio_measurements([
            ["BIA", "POBRANIE", "A", "2025-03-28", 1],
            ["BIA", "POBRANIE", "B", "2025-03-29", 1],
            ["BIA", "POBRANIE", "B", "2025-03-31", 1],
            ["BIA", "POBRANIE", "A", "2025-04-01", 1],
        ])
        reports = analyze_portfolio(data, grace_days=3, ignored_dates=["2025-03-30"])
        gap = reports["przerwy_pomiarow_mdd"].iloc[0]
        self.assertEqual(gap.klient_nazwa, "A")
        self.assertEqual(gap.dni_przerwy, 2)
        clients = reports["portfel_mdd"].set_index("klient_nazwa")
        self.assertEqual(clients.loc["B", "dni_bez_pomiarow_na_koncu"], 1)
        self.assertEqual(reports["zmiany_portfela_dzienne"].status.iloc[-1], "ZMIANA_ZESTAWU_MDD")

    def test_all_rows_excluded_gives_clear_error_and_report_without_querying_weather(self):
        with TemporaryDirectory() as tmp, redirect_stdout(StringIO()):
            source = Path(tmp) / "energia.csv"
            pd.DataFrame([energy_row("2025-03-30", 3)], columns=ENERGY_COLUMNS).to_csv(source, index=False)
            directory = Path(tmp) / "model"
            with patch("mdd_forecasting.branch_cli._weather", side_effect=AssertionError("Bez pogody")):
                with self.assertRaisesRegex(ValueError, "Brak pomiarów po filtrze"):
                    train_main(["--energy", str(source), "--artifacts-dir", str(directory), "--weather", "unused.csv"])
            report = pd.read_csv(directory / "pominiete_dni_zmiany_czasu.csv", sep=";", decimal=",")
            self.assertEqual(report.rekordy_wykluczone.tolist(), [1])

    def test_cli_include_option_restores_original_validation(self):
        with TemporaryDirectory() as tmp, redirect_stdout(StringIO()):
            source = Path(tmp) / "energia.csv"
            pd.DataFrame([energy_row("2025-03-30", 3)], columns=ENERGY_COLUMNS).to_csv(source, index=False)
            with self.assertRaises(EnergyTimeError):
                train_main(["--energy", str(source), "--artifacts-dir", tmp, "--weather", "unused.csv", "--include-dst-days"])

    def test_cli_training_saves_rule_and_excludes_dst_from_holdout(self):
        times = pd.date_range("2025-02-15T00:00Z", "2025-04-01T23:00Z", freq="h")
        rows, weather_rows = [], []
        for index, timestamp in enumerate(times):
            local = timestamp.tz_convert("Europe/Warsaw")
            date = local.strftime("%Y-%m-%d")
            hour = (local.hour or 24) if date == "2025-03-30" else local.hour + 1
            for direction, value in [(1, 100 + index % 24), (-1, 20 if 8 <= local.hour < 18 else 0)]:
                rows.append(energy_row(date, hour, direction=direction, value=value))
            issue = timestamp - pd.Timedelta(hours=96)
            weather_rows.append({"punkt": "Białystok", "dataGodzinaUTC": timestamp,
                                 "dataGodzinaCET": local.tz_localize(None),
                                 "czasDanychZrodlaUTC": issue,
                                 "czasDanychZrodlaCET": issue.tz_convert("Europe/Warsaw").tz_localize(None),
                                 "temperatura": 10 + np.sin(index / 24)})
        with TemporaryDirectory() as tmp, redirect_stdout(StringIO()):
            directory = Path(tmp)
            energy, weather = directory / "energia.csv", directory / "pogoda.csv"
            pd.DataFrame(rows, columns=ENERGY_COLUMNS).to_csv(energy, index=False)
            pd.DataFrame(weather_rows).to_csv(weather, index=False)
            model = directory / "model"
            result = train_main(["--energy", str(energy), "--artifacts-dir", str(model), "--weather", str(weather),
                                 "--max-iter", "2", "--validation-days", "4", "--test-days", "4",
                                 "--label-delay-hours", "24", "--importance-rows", "0", "--threads", "2",
                                 "--weather-variants", "temperatura", "--min-evaluation-targets", "1"])
            self.assertEqual(result, 0)
            bundle = load_bundle(model)
            self.assertTrue(bundle["input_schema"]["exclude_dst_days"])
            manifest = json.loads((model / "manifest_oddzialy.json").read_text(encoding="utf-8"))
            self.assertTrue(manifest["input_schema"]["exclude_dst_days"])
            history = pd.read_csv(model / "historia_oddzialow.csv", sep=";")
            self.assertFalse(history.doba_handlowa.eq("2025-03-30").any())
            heldout = pd.read_csv(model / "backtest_bilans.csv", sep=";")
            targets = pd.to_datetime(heldout.target_time_utc, utc=True).dt.tz_convert("Europe/Warsaw")
            self.assertFalse(targets.dt.strftime("%Y-%m-%d").eq("2025-03-30").any())
            excluded = pd.read_csv(model / "pominiete_dni_zmiany_czasu.csv", sep=";", decimal=",")
            self.assertEqual(excluded.rekordy_wykluczone.tolist(), [46])
            from openpyxl import load_workbook
            workbook = load_workbook(model / "raport_treningu.xlsx", read_only=True)
            self.assertIn("Pominiete_doby", workbook.sheetnames)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
